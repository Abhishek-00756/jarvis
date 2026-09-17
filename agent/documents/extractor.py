"""Document reading: PDF, DOCX, PPTX, spreadsheets."""

from pathlib import Path

MAX_CHARS = 8000


def read_pdf(path: str) -> str:
    from pypdf import PdfReader
    p = Path(path).expanduser()
    reader = PdfReader(str(p))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return text[:MAX_CHARS]


def search_pdf(path: str, query: str) -> str:
    from pypdf import PdfReader
    p = Path(path).expanduser()
    reader = PdfReader(str(p))
    matches = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if query.lower() in text.lower():
            idx = text.lower().find(query.lower())
            snippet = text[max(0, idx - 100): idx + 200]
            matches.append(f"Page {i + 1}: ...{snippet}...")
    return "\n\n".join(matches) if matches else f"'{query}' not found in {path}."


def read_docx(path: str) -> str:
    import docx
    p = Path(path).expanduser()
    document = docx.Document(str(p))
    text = "\n".join(para.text for para in document.paragraphs)
    return text[:MAX_CHARS]


def read_pptx(path: str) -> str:
    from pptx import Presentation
    p = Path(path).expanduser()
    prs = Presentation(str(p))
    slides_text = []
    for i, slide in enumerate(prs.slides):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text]
        slides_text.append(f"Slide {i + 1}: " + " | ".join(texts))
    return "\n".join(slides_text)[:MAX_CHARS]


def read_spreadsheet(path: str, sheet_name: str = "") -> str:
    import openpyxl
    p = Path(path).expanduser()
    if p.suffix.lower() == ".csv":
        import csv
        with open(p) as f:
            rows = list(csv.reader(f))
        return "\n".join(", ".join(row) for row in rows[:50])[:MAX_CHARS]
    wb = openpyxl.load_workbook(str(p), data_only=True)
    sheet = wb[sheet_name] if sheet_name else wb.active
    rows = []
    for row in sheet.iter_rows(max_row=50, values_only=True):
        rows.append(", ".join(str(c) if c is not None else "" for c in row))
    return "\n".join(rows)[:MAX_CHARS]
